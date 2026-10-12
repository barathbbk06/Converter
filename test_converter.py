"""
Test Suite for Local Word <-> PDF Converter
-------------------------------------------
Tests:
1. Sample .docx generation.
2. Word to PDF conversion (Word COM / LibreOffice).
3. PDF to Word conversion (pdf2docx).
4. Duplicate output collision handling (unique filename generator).
5. Error handling for non-existent and unsupported files.
"""

import os
import sys
import shutil
from pathlib import Path
import docx

# Ensure UTF-8 output on Windows console
if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

from converter import (
    convert_word_to_pdf,
    convert_pdf_to_word,
    get_unique_filepath,
    get_default_output_path,
    is_ms_word_available,
    get_libreoffice_path,
    ConversionError
)


def run_tests():
    print("=" * 60)
    print("Running Local Word <-> PDF Converter Test Suite")
    print("=" * 60)

    # 1. Environment Verification
    word_ok, word_info = is_ms_word_available()
    lo_path = get_libreoffice_path()
    print(f"[*] Microsoft Word Status : {word_info if word_ok else 'Not Available (' + word_info + ')'}")
    print(f"[*] LibreOffice Status    : {lo_path if lo_path else 'Not Detected'}")

    if not word_ok and not lo_path:
        print("[!] WARNING: Neither Microsoft Word nor LibreOffice is available for Word->PDF.")

    # 2. Setup Test Workspace
    test_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), "test_run_artifacts"))
    os.makedirs(test_dir, exist_ok=True)
    print(f"[*] Test directory created: {test_dir}")

    try:
        # Step A: Create sample docx
        sample_docx = os.path.join(test_dir, "Sample Report.docx")
        doc = docx.Document()
        doc.add_heading("Local Document Conversion Test Report", level=1)
        doc.add_paragraph("This is an automated test document verifying local conversion between Word and PDF.")
        doc.add_paragraph("Formatting check: Bold text, italic text, and multiple paragraphs.")

        # Add a table
        table = doc.add_table(rows=3, cols=3)
        headers = ["ID", "Feature", "Status"]
        for col_idx, h in enumerate(headers):
            table.cell(0, col_idx).text = h

        rows_data = [
            ("1", "Offline Word to PDF", "Supported"),
            ("2", "Offline PDF to Word", "Supported")
        ]
        for row_idx, r in enumerate(rows_data, start=1):
            for col_idx, val in enumerate(r):
                table.cell(row_idx, col_idx).text = val

        doc.save(sample_docx)
        assert os.path.exists(sample_docx), "Failed to create sample DOCX"
        print(f"[OK] Step A: Created sample DOCX ({os.path.getsize(sample_docx)} bytes)")

        # Step B: Test Word to PDF conversion
        sample_pdf = get_default_output_path(sample_docx, test_dir, "pdf")
        converted_pdf = convert_word_to_pdf(sample_docx, sample_pdf)
        assert os.path.exists(converted_pdf), "PDF conversion output does not exist"
        assert os.path.getsize(converted_pdf) > 0, "PDF conversion produced an empty file"
        print(f"[OK] Step B: Word -> PDF conversion succeeded ({os.path.getsize(converted_pdf)} bytes)")

        # Step C: Test Collision Handling
        dup_candidate = get_unique_filepath(converted_pdf)
        assert dup_candidate != converted_pdf, "Collision did not produce unique filename"
        assert "(1)" in os.path.basename(dup_candidate), f"Expected '(1)' in {dup_candidate}"
        print(f"[OK] Step C: Collision handler generated unique path: {os.path.basename(dup_candidate)}")

        # Step D: Test PDF to Word conversion
        sample_out_docx = os.path.join(test_dir, "Converted Back From PDF.docx")
        converted_word = convert_pdf_to_word(converted_pdf, sample_out_docx)
        assert os.path.exists(converted_word), "PDF to Word output does not exist"
        assert os.path.getsize(converted_word) > 0, "PDF to Word produced an empty file"

        # Verify converted Word document content
        res_doc = docx.Document(converted_word)
        full_text = " ".join([p.text for p in res_doc.paragraphs])
        assert "Report" in full_text or "Local" in full_text, "Extracted text content verification failed"
        print(f"[OK] Step D: PDF -> Word conversion succeeded ({os.path.getsize(converted_word)} bytes)")

        # Step E: Test Error Handling
        print("[*] Step E: Testing edge cases & error handling...")

        # Non-existent file
        try:
            convert_word_to_pdf(os.path.join(test_dir, "non_existent.docx"), os.path.join(test_dir, "out.pdf"))
            assert False, "Failed to raise ConversionError on missing file"
        except ConversionError:
            print("  [OK] Handled non-existent file correctly.")

        # Invalid extension for Word->PDF
        try:
            invalid_file = os.path.join(test_dir, "invalid.txt")
            with open(invalid_file, "w") as f:
                f.write("test")
            convert_word_to_pdf(invalid_file, os.path.join(test_dir, "out.pdf"))
            assert False, "Failed to raise ConversionError on unsupported extension"
        except ConversionError:
            print("  [OK] Handled unsupported extension correctly.")

        print("=" * 60)
        print("ALL TESTS PASSED SUCCESSFULLY! (100% Offline & Fully Functional)")
        print("=" * 60)

    finally:
        # Cleanup test directory
        if os.path.exists(test_dir):
            shutil.rmtree(test_dir, ignore_errors=True)
            print("[*] Test directory cleaned up.")


if __name__ == "__main__":
    run_tests()
